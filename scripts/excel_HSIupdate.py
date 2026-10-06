"""Apply HSI_min updates in source Excel files from a summary workbook."""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


DEFAULT_FOLDER = Path(r"E:/2026/ZoemZoneLimburg/OutputDataInfo")


@dataclass
class UpdateInstruction:
    row_number: int
    id1: str
    id2: str
    old_hsi: Any
    new_hsi: Any
    source_excel: str


def _cell_to_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _normalize_header(value: Any) -> str:
    return _cell_to_text(value).lower()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Update HSI_min values in source Excel files from a summary workbook.",
    )
    parser.add_argument(
        "--folder",
        default=str(DEFAULT_FOLDER),
        help="Folder that contains summary and source Excel files.",
    )
    parser.add_argument(
        "--summary",
        default=None,
        help="Path to summary workbook. If omitted, newest summary_*.xlsx in folder is used.",
    )
    return parser


def _find_summary_file(folder: Path) -> Path:
    candidates = sorted(
        [
            path
            for path in folder.glob("summary_*.xlsx")
            if path.is_file() and not path.name.startswith("~$")
        ],
        key=lambda item: item.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError(f"No summary_*.xlsx file found in {folder}")
    return candidates[0]


def _read_instructions(summary_path: Path) -> list[UpdateInstruction]:
    workbook = load_workbook(filename=summary_path, data_only=True, read_only=True)
    sheet = workbook.active
    if sheet is None:
        workbook.close()
        raise RuntimeError("Summary workbook has no active worksheet.")

    instructions: list[UpdateInstruction] = []

    for excel_row in range(3, sheet.max_row + 1):
        id1 = _cell_to_text(sheet.cell(row=excel_row, column=1).value)
        id2 = _cell_to_text(sheet.cell(row=excel_row, column=2).value)
        old_hsi = sheet.cell(row=excel_row, column=3).value
        new_hsi = sheet.cell(row=excel_row, column=4).value
        source_excel = _cell_to_text(sheet.cell(row=excel_row, column=5).value)

        if not id1 and not source_excel and new_hsi in (None, ""):
            continue
        if new_hsi in (None, ""):
            continue
        if not source_excel:
            continue

        instructions.append(
            UpdateInstruction(
                row_number=excel_row,
                id1=id1,
                id2=id2,
                old_hsi=old_hsi,
                new_hsi=new_hsi,
                source_excel=source_excel,
            )
        )

    workbook.close()
    return instructions


def _find_hsi_col_and_id2_logic(sheet: Any) -> tuple[int, bool]:
    header_row = 1
    max_probe_rows = min(sheet.max_row, 20)
    for row_index in range(1, max_probe_rows + 1):
        headers = [_normalize_header(sheet.cell(row=row_index, column=col).value) for col in range(1, sheet.max_column + 1)]
        if "hsi_min" in headers:
            header_row = row_index
            break

    headers = [_normalize_header(sheet.cell(row=header_row, column=col).value) for col in range(1, sheet.max_column + 1)]
    if "hsi_min" not in headers:
        raise ValueError("Could not find HSI_min column")

    hsi_col = headers.index("hsi_min") + 1
    second_is_count = len(headers) >= 2 and headers[1] == "count"
    return hsi_col, second_is_count


def _apply_updates_to_source(
    source_path: Path,
    updates: list[UpdateInstruction],
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []

    workbook = load_workbook(filename=source_path)
    updated_any = False

    try:
        for instruction in updates:
            matched = False
            applied = 0

            for sheet in workbook.worksheets:
                try:
                    hsi_col, second_is_count = _find_hsi_col_and_id2_logic(sheet)
                except ValueError:
                    continue

                header_row = 1
                max_probe_rows = min(sheet.max_row, 20)
                for row_index in range(1, max_probe_rows + 1):
                    headers = [_normalize_header(sheet.cell(row=row_index, column=col).value) for col in range(1, sheet.max_column + 1)]
                    if "hsi_min" in headers:
                        header_row = row_index
                        break

                for row_index in range(header_row + 1, sheet.max_row + 1):
                    row_id1 = _cell_to_text(sheet.cell(row=row_index, column=1).value)
                    if not row_id1 or row_id1 != instruction.id1:
                        continue

                    if second_is_count:
                        if instruction.id2:
                            continue
                    else:
                        row_id2 = _cell_to_text(sheet.cell(row=row_index, column=2).value)
                        if instruction.id2:
                            if row_id2 != instruction.id2:
                                continue
                        elif row_id2:
                            continue

                    matched = True
                    current_value = sheet.cell(row=row_index, column=hsi_col).value
                    if current_value != instruction.new_hsi:
                        sheet.cell(row=row_index, column=hsi_col).value = instruction.new_hsi
                        updated_any = True
                        applied += 1

            status = "updated" if applied > 0 else ("already_same" if matched else "not_found")
            results.append(
                {
                    "summary_row": instruction.row_number,
                    "source_excel": source_path.name,
                    "id1": instruction.id1,
                    "id2": instruction.id2,
                    "old_hsi_summary": instruction.old_hsi,
                    "new_hsi": instruction.new_hsi,
                    "status": status,
                    "matched_rows": applied if applied > 0 else (1 if matched else 0),
                }
            )
    finally:
        if updated_any:
            workbook.save(source_path)
        workbook.close()

    return results


def _write_report(output_path: Path, run_timestamp: str, summary_file: str, rows: list[dict[str, Any]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("Could not create report worksheet")
    sheet.title = "HSI_updates"

    sheet.append(["run_timestamp", run_timestamp, "summary_file", summary_file])
    sheet.append(
        [
            "summary_row",
            "source_excel",
            "id1",
            "id2",
            "old_hsi_summary",
            "new_hsi",
            "status",
            "matched_rows",
        ]
    )

    for row in rows:
        sheet.append(
            [
                row["summary_row"],
                row["source_excel"],
                row["id1"],
                row["id2"],
                row["old_hsi_summary"],
                row["new_hsi"],
                row["status"],
                row["matched_rows"],
            ]
        )

    # Autosize all report columns based on content length.
    for col_idx, column_cells in enumerate(sheet.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        col_letter = get_column_letter(col_idx)
        sheet.column_dimensions[col_letter].width = max(10, min(max_len + 2, 80))

    source_sheet = workbook.create_sheet(title="source_excel_overzicht")
    source_sheet.append(
        [
            "source_excel",
            "updated_rows",
            "already_same_rows",
            "total_rows_with_changes",
        ]
    )

    per_file_counts: dict[str, dict[str, int]] = {}
    for row in rows:
        source_excel = str(row.get("source_excel", "")).strip()
        if not source_excel:
            continue

        status = str(row.get("status", "")).strip().lower()
        if status not in {"updated", "already_same"}:
            continue

        counts = per_file_counts.setdefault(
            source_excel,
            {"updated": 0, "already_same": 0},
        )
        if status == "updated":
            counts["updated"] += 1
        elif status == "already_same":
            counts["already_same"] += 1

    for source_excel in sorted(per_file_counts):
        counts = per_file_counts[source_excel]
        source_sheet.append(
            [
                source_excel,
                counts["updated"],
                counts["already_same"],
                counts["updated"] + counts["already_same"],
            ]
        )

    for col_idx, column_cells in enumerate(source_sheet.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_len:
                max_len = len(value)
        col_letter = get_column_letter(col_idx)
        source_sheet.column_dimensions[col_letter].width = max(10, min(max_len + 2, 80))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    folder = Path(str(args.folder).strip()).resolve()
    if not folder.exists() or not folder.is_dir():
        parser.error(f"Folder does not exist or is not a directory: {folder}")

    summary_path = Path(args.summary).resolve() if args.summary else _find_summary_file(folder)
    if not summary_path.exists():
        parser.error(f"Summary file not found: {summary_path}")

    instructions = _read_instructions(summary_path)
    if not instructions:
        parser.error("No updates found: column 4 contains no new HSI values.")

    by_source: dict[str, list[UpdateInstruction]] = {}
    for instruction in instructions:
        by_source.setdefault(instruction.source_excel, []).append(instruction)

    all_results: list[dict[str, Any]] = []

    for source_name, source_updates in sorted(by_source.items()):
        source_path = folder / source_name
        if not source_path.exists():
            for instruction in source_updates:
                all_results.append(
                    {
                        "summary_row": instruction.row_number,
                        "source_excel": source_name,
                        "id1": instruction.id1,
                        "id2": instruction.id2,
                        "old_hsi_summary": instruction.old_hsi,
                        "new_hsi": instruction.new_hsi,
                        "status": "source_missing",
                        "matched_rows": 0,
                    }
                )
            continue

        all_results.extend(_apply_updates_to_source(source_path, source_updates))

    run_timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = folder / f"HSI_updates_{run_timestamp}.xlsx"
    _write_report(
        output_path=output_path,
        run_timestamp=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        summary_file=summary_path.name,
        rows=all_results,
    )

    updated_count = sum(1 for row in all_results if row["status"] == "updated")
    skipped_count = len(all_results) - updated_count
    print(f"Summary input: {summary_path}")
    print(f"Instructions processed: {len(all_results)}")
    print(f"Updated rows: {updated_count}")
    print(f"Other statuses: {skipped_count}")
    print(f"Report written: {output_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
